import pytest
import io
import os
from reportlab.pdfgen import canvas
import docx
import openpyxl
from app.db.session import SessionLocal
from app.models.user import User
from app.models.project import Project
from app.services.document_processor import (
    sanitize_filename,
    validate_file_content,
    DocumentValidationError
)

@pytest.fixture
def auth_header(client):
    """Obtain a fresh auth token via the login endpoint."""
    from app.core.rate_limit import RateLimitMiddleware
    RateLimitMiddleware.reset()
    resp = client.post("/api/v1/auth/login", json={
        "email": "officer@cpwd.gov.in",
        "password": os.environ["TEST_DEMO_PASSWORD"]
    })
    assert resp.status_code == 200, f"Login failed: {resp.json()}"
    token = resp.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}

@pytest.fixture
def test_project():
    db = SessionLocal()
    user = db.query(User).filter(User.email == "officer@cpwd.gov.in").first()
    proj = db.query(Project).filter(Project.code == "PROJ-TEST-001").first()
    if not proj:
        proj = Project(
            code="PROJ-TEST-001",
            title="Bridge Construction CPWD",
            description="Testing tender document processing",
            created_by_id=user.id
        )
        db.add(proj)
        db.commit()
        db.refresh(proj)
    proj_id = proj.id
    db.close()
    return proj_id

def create_sample_pdf(text_lines: list[str]) -> bytes:
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer)
    y = 800
    for line in text_lines:
        c.drawString(50, y, line)
        y -= 25
        if y < 50:
            c.showPage()
            y = 800
    c.save()
    buffer.seek(0)
    return buffer.getvalue()

def create_sample_docx(text_paragraphs: list[str]) -> bytes:
    doc = docx.Document()
    for p in text_paragraphs:
        doc.add_paragraph(p)
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def create_sample_xlsx(rows: list[list[str]]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Specifications"
    for r in rows:
        ws.append(r)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

# 1. Normal PDF test
def test_upload_normal_pdf(client, auth_header, test_project):
    content = create_sample_pdf([
        "Tender Specification Document: Reinforced Concrete Works",
        "All cement concrete shall strictly comply with IS 456:2000 specifications.",
        "Steel reinforcement bars shall conform to IS 1786 Fe 500D grade with BIS mark."
    ])
    files = {"file": ("tender_spec.pdf", content, "application/pdf")}
    data = {"project_id": test_project}
    
    response = client.post("/api/v1/documents/upload", files=files, data=data, headers=auth_header)
    assert response.status_code == 201
    res_data = response.json()
    assert res_data["filename"] == "tender_spec.pdf"
    assert res_data["file_type"] == "pdf"
    assert res_data["processing_status"] == "completed"
    assert res_data["ocr_applied"] is False
    assert res_data["total_pages"] >= 1

# 2. Scanned PDF test (sparse text triggers OCR)
def test_upload_scanned_pdf(client, auth_header, test_project):
    content = create_sample_pdf(["A"]) # Only 1 character -> triggers OCR fallback
    files = {"file": ("scanned_contract.pdf", content, "application/pdf")}
    data = {"project_id": test_project}

    response = client.post("/api/v1/documents/upload", files=files, data=data, headers=auth_header)
    assert response.status_code == 201
    res_data = response.json()
    assert res_data["ocr_applied"] is True
    assert res_data["processing_status"] == "completed"

# 3. Large PDF test (multiple pages & chunking)
def test_upload_large_pdf(client, auth_header, test_project):
    lines = [f"Clause {i}: Technical specification parameter details and testing requirements for steel and concrete works under IS 456 and IS 1786 standards." for i in range(1, 100)]
    content = create_sample_pdf(lines)
    files = {"file": ("large_tender.pdf", content, "application/pdf")}
    data = {"project_id": test_project}

    response = client.post("/api/v1/documents/upload", files=files, data=data, headers=auth_header)
    assert response.status_code == 201
    doc_id = response.json()["id"]

    # Check chunks
    chunks_resp = client.get(f"/api/v1/documents/{doc_id}/chunks", headers=auth_header)
    assert chunks_resp.status_code == 200
    chunks = chunks_resp.json()
    assert len(chunks) > 1
    assert chunks[0]["page_number"] >= 1

# 4. Malformed PDF test
def test_upload_malformed_pdf(client, auth_header, test_project):
    corrupted_content = b"Not a real PDF file header but has .pdf extension"
    files = {"file": ("corrupted.pdf", corrupted_content, "application/pdf")}
    data = {"project_id": test_project}

    response = client.post("/api/v1/documents/upload", files=files, data=data, headers=auth_header)
    assert response.status_code == 400
    assert "Magic bytes mismatch" in response.json()["detail"]

# 5. Unsupported file type test
def test_upload_unsupported_file(client, auth_header, test_project):
    files = {"file": ("payload.exe", b"MZ\x90\x00\x03\x00\x00\x00", "application/octet-stream")}
    data = {"project_id": test_project}

    response = client.post("/api/v1/documents/upload", files=files, data=data, headers=auth_header)
    assert response.status_code == 415
    assert "Unsupported format" in response.json()["detail"]

# 6. Malicious filename sanitization test
def test_upload_malicious_filename(client, auth_header, test_project):
    content = b"Simple valid plain text document for testing tender specs."
    files = {"file": ("../../../../etc/passwd_spec.txt", content, "text/plain")}
    data = {"project_id": test_project}

    response = client.post("/api/v1/documents/upload", files=files, data=data, headers=auth_header)
    assert response.status_code == 201
    filename = response.json()["filename"]
    # Path traversal should be stripped
    assert ".." not in filename
    assert "/" not in filename
    assert "\\" not in filename
    assert filename == "passwd_spec.txt"

# 7. Oversized file test
def test_upload_oversized_file_validation():
    oversized_bytes = b"0" * (51 * 1024 * 1024) # 51 MB
    with pytest.raises(DocumentValidationError) as exc_info:
        validate_file_content("huge.txt", oversized_bytes, max_size_mb=50)
    assert "exceeds maximum limit" in str(exc_info.value)

# 8. DOCX upload test
def test_upload_docx(client, auth_header, test_project):
    content = create_sample_docx([
        "Procurement Specification for Electrical Cables",
        "Cables must strictly comply with IS 694 for PVC insulation.",
        "Earthing design shall adhere to IS 3043."
    ])
    files = {"file": ("cables_spec.docx", content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    data = {"project_id": test_project}

    response = client.post("/api/v1/documents/upload", files=files, data=data, headers=auth_header)
    assert response.status_code == 201
    assert response.json()["file_type"] == "docx"

# 9. XLSX upload test
def test_upload_xlsx(client, auth_header, test_project):
    content = create_sample_xlsx([
        ["Item No", "Description", "Standard Code", "Grade"],
        ["1", "Deformed Steel Bars", "IS 1786", "Fe 500D"],
        ["2", "Structural Steel", "IS 2062", "E250"]
    ])
    files = {"file": ("bill_of_quantities.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    data = {"project_id": test_project}

    response = client.post("/api/v1/documents/upload", files=files, data=data, headers=auth_header)
    assert response.status_code == 201
    assert response.json()["file_type"] == "xlsx"

# 10. Filename sanitization unit test
def test_sanitize_filename_edge_cases():
    assert sanitize_filename("..\\..\\windows\\system32\\cmd.exe") == "cmd.exe"
    assert sanitize_filename("../../../etc/shadow") == "shadow"
    sanitized_cmd = sanitize_filename("tender; rm -rf /; .pdf")
    assert "tender" in sanitized_cmd
    assert ";" not in sanitized_cmd
    assert "/" not in sanitized_cmd
    assert sanitized_cmd.endswith(".pdf")
    assert sanitize_filename("normal_file.pdf") == "normal_file.pdf"
